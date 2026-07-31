from __future__ import annotations

import hashlib
import io
import json
import re
import shutil
import unicodedata
import zipfile
from datetime import datetime, timedelta, timezone
from pathlib import Path, PurePosixPath
from uuid import uuid4
from urllib.parse import urlparse

from email_validator import EmailNotValidError, validate_email
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.services.auditoria_service import registrar_bitacora
from app.services.identidad_importacion_service import (
    buscar_id_empresa_reutilizable,
    marcar_id_reutilizado,
)
from app.services.upload_security import normalizar_nombre_archivo, validar_documento_usuario
from infrastructure.database.dependencies import obtener_db
from infrastructure.persistence.models.documento_empresa import DocumentoEmpresaModel
from infrastructure.persistence.models.empresa import EmpresaModel
from infrastructure.persistence.models.responsable_empresa import ResponsableEmpresaModel
from infrastructure.persistence.models.tipo_unidad_receptora import TipoUnidadReceptoraModel
from infrastructure.persistence.models.solicitud_empresa import SolicitudEmpresaModel
from infrastructure.persistence.models.tipo_documento_empresa import TipoDocumentoEmpresaModel
from infrastructure.persistence.models.usuario import UsuarioModel
from infrastructure.security.auth_dependencies import obtener_usuario_actual, requerir_roles

router = APIRouter(
    prefix="/coord-unidades/empresas/importacion",
    tags=["Coordinador Unidades - Importacion de Empresas"],
    dependencies=[Depends(requerir_roles(["Coordinador de Unidades Receptoras", "Administrador"]))],
)

UPLOADS_ROOT = Path(__file__).resolve().parents[3] / "uploads"
IMPORTACIONES_ROOT = UPLOADS_ROOT / "importaciones_empresas"
DOCUMENTOS_ROOT = UPLOADS_ROOT / "empresas" / "importacion_masiva"
MAX_EXCEL_BYTES = 5 * 1024 * 1024
MAX_ZIP_BYTES = 50 * 1024 * 1024
MAX_PDF_BYTES = 2 * 1024 * 1024
MAX_FILAS = 500
MAX_ARCHIVOS_ZIP = 250
VIGENCIA_JOB_HORAS = 24
HOJA_DATOS = "Unidades receptoras"
HOJA_CATALOGOS = "Catálogos"
HOJA_INSTRUCCIONES = "Instrucciones"
FILA_ENCABEZADOS = 9
FILA_INICIO_DATOS = 10
FILA_FIN_CAPTURA = 30
FILA_TITULO_EJEMPLO = 32
FILA_EJEMPLO = 33

COLUMNA_DOCUMENTOS = (
    "Documentos en formato PDF:\n"
    "• Acta constitutiva con datos de registro para personas jurídicas.\n"
    "• Acta de nacimiento para persona física.\n"
    "• Acta de la sesión de Cabildo para ayuntamientos.\n"
    "• Acuerdo de Cabildo de creación del órgano para organismos descentralizados o desconcentrados.\n"
    "• Decreto de creación del ente público para el sector público estatal o federal.\n"
    "• Constancia de situación fiscal.\n"
    "• Comprobante de domicilio expedido con una antigüedad máxima de tres meses.\n"
    "• Identificación oficial con fotografía de la persona apoderada."
)
COLUMNAS = [
    "Nombre oficial de la institución, empresa u organización",
    "Tipo de unidad receptora",
    "Registro Federal de Contribuyentes (RFC)",
    "Domicilio completo (calle, número, colonia y código postal)",
    "Teléfono(s) de contacto institucional",
    "Correo electrónico oficial",
    "Horario de atención de la institución",
    "Nombre completo",
    "Cargo o puesto dentro de la institución",
    "Área o departamento al que pertenece",
    "Teléfono de contacto",
    "Correo electrónico",
    "Áreas o departamentos en los que se recibirán estudiantes en prácticas",
    "Número de estudiantes que puede recibir por periodo",
    "Perfil académico requerido (programa educativo, conocimientos o habilidades)",
    "Actividades principales que realizará el estudiante durante sus prácticas",
    "Horario propuesto para la realización de las prácticas",
    "Modalidad de las prácticas",
    COLUMNA_DOCUMENTOS,
    "Municipio",
    "Estado",
    "Estatus",
    "Observaciones",
    "Carta de colaboración",
]
CAMPOS_OBLIGATORIOS = {COLUMNAS[0]}
# Se conserva vacío únicamente para no romper compatibilidad interna con trabajos validados previamente.
COLUMNAS_DOCUMENTOS: dict[str, str] = {}
TIPOS_UNIDAD = [
    "Sector Productivo - Persona física",
    "Sector Productivo - Persona jurídica",
    "Sector Público - Estatal o Federal",
    "Sector Social",
    "Sector Municipal - Ayuntamiento",
    "Sector Municipal - Descentralizado o desconcentrado",
]
MODALIDADES = ["Presencial", "Virtual", "Híbrida"]
ESTATUS_PERMITIDOS = ["Solicitante", "Pendiente", "Aceptado"]
CARTA_COLABORACION = ["Sí", "No", "Pendiente"]
ESTADOS_MEXICO = [
    "Aguascalientes", "Baja California", "Baja California Sur", "Campeche", "Chiapas", "Chihuahua",
    "Ciudad de México", "Coahuila", "Colima", "Durango", "Estado de México", "Guanajuato",
    "Guerrero", "Hidalgo", "Jalisco", "Michoacán", "Morelos", "Nayarit", "Nuevo León",
    "Oaxaca", "Puebla", "Querétaro", "Quintana Roo", "San Luis Potosí", "Sinaloa", "Sonora",
    "Tabasco", "Tamaulipas", "Tlaxcala", "Veracruz", "Yucatán", "Zacatecas",
]
RFC_PATTERN = re.compile(r"^[A-ZÑ&]{3,4}\d{6}[A-Z0-9]{3}$")
TELEFONO_PATTERN = re.compile(r"^\+?\d{7,15}$")
MUNICIPIO_PATTERN = re.compile(r"^[A-Za-zÁÉÍÓÚÜÑáéíóúüñ .'-]{2,100}$")


def _texto(valor: object) -> str:
    if valor is None:
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor).strip()


def _clave_catalogo(valor: str) -> str:
    return unicodedata.normalize("NFKD", valor).encode("ascii", "ignore").decode("ascii").casefold().strip()


def _catalogo(valor: str, opciones: list[str]) -> str | None:
    clave = _clave_catalogo(valor)
    return next((item for item in opciones if _clave_catalogo(item) == clave), None)


def _normalizar_rfc(valor: str) -> str:
    return re.sub(r"[\s-]+", "", valor).upper()


def _normalizar_telefono(valor: str) -> str:
    return re.sub(r"[\s\-()]+", "", valor)


def _correo_valido(valor: str) -> str | None:
    try:
        return validate_email(valor, check_deliverability=False).normalized.lower()
    except EmailNotValidError:
        return None


def _enlace_documentos_valido(valor: str) -> bool:
    if not valor or len(valor) > 1000:
        return False
    try:
        enlace = urlparse(valor)
    except ValueError:
        return False
    return enlace.scheme.lower() in {"http", "https"} and bool(enlace.netloc)


def _normalizar_telefonos(valor: str) -> str | None:
    telefonos = [item.strip() for item in re.split(r"[,;/]+", valor) if item.strip()]
    if not telefonos:
        return None
    normalizados = [_normalizar_telefono(item) for item in telefonos]
    if any(not TELEFONO_PATTERN.fullmatch(item) for item in normalizados):
        return None
    return "; ".join(normalizados)


def _nombre_pdf(valor: str) -> str | None:
    if not valor or len(valor) > 120:
        return None
    ruta = PurePosixPath(valor.replace("\\", "/"))
    if ruta.is_absolute() or len(ruta.parts) != 1 or ruta.name in {"", ".", ".."} or ruta.suffix.lower() != ".pdf":
        return None
    try:
        return normalizar_nombre_archivo(ruta.name, "documento.pdf")
    except HTTPException:
        return None


def _sha256(contenido: bytes) -> str:
    return hashlib.sha256(contenido).hexdigest()


def _limpiar_jobs_expirados() -> None:
    IMPORTACIONES_ROOT.mkdir(parents=True, exist_ok=True)
    limite = datetime.now(timezone.utc) - timedelta(hours=VIGENCIA_JOB_HORAS)
    for carpeta in IMPORTACIONES_ROOT.iterdir():
        if not carpeta.is_dir():
            continue
        try:
            if datetime.fromtimestamp(carpeta.stat().st_mtime, timezone.utc) < limite:
                shutil.rmtree(carpeta, ignore_errors=True)
        except OSError:
            continue


def _validar_excel_upload(archivo: UploadFile, contenido: bytes) -> None:
    if Path(archivo.filename or "").suffix.lower() != ".xlsx":
        raise HTTPException(status_code=400, detail="El archivo de empresas debe tener extensión .xlsx")
    if len(contenido) > MAX_EXCEL_BYTES:
        raise HTTPException(status_code=400, detail="El archivo Excel excede el límite de 5 MB")
    if not contenido.startswith(b"PK"):
        raise HTTPException(status_code=400, detail="El contenido no corresponde a un XLSX válido")
    mime = (archivo.content_type or "").split(";", 1)[0].lower()
    if mime not in {"", "application/octet-stream", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}:
        raise HTTPException(status_code=400, detail="Tipo MIME de Excel no permitido")


def _validar_zip_upload(archivo: UploadFile | None, contenido: bytes | None) -> None:
    if archivo is None or contenido is None:
        return
    if Path(archivo.filename or "").suffix.lower() != ".zip":
        raise HTTPException(status_code=400, detail="Los documentos deben enviarse en un archivo .zip")
    if len(contenido) > MAX_ZIP_BYTES:
        raise HTTPException(status_code=400, detail="El archivo ZIP excede el límite de 50 MB")
    if not contenido.startswith(b"PK"):
        raise HTTPException(status_code=400, detail="El contenido no corresponde a un ZIP válido")
    mime = (archivo.content_type or "").split(";", 1)[0].lower()
    if mime not in {"", "application/zip", "application/x-zip-compressed", "application/octet-stream"}:
        raise HTTPException(status_code=400, detail="Tipo MIME de ZIP no permitido")


def _indice_zip(contenido: bytes | None) -> tuple[dict[str, zipfile.ZipInfo], list[str]]:
    if contenido is None:
        return {}, []
    errores: list[str] = []
    indice: dict[str, zipfile.ZipInfo] = {}
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as archivo_zip:
            infos = [item for item in archivo_zip.infolist() if not item.is_dir()]
            if len(infos) > MAX_ARCHIVOS_ZIP:
                return {}, [f"El ZIP contiene más de {MAX_ARCHIVOS_ZIP} archivos"]
            for info in infos:
                ruta = PurePosixPath(info.filename.replace("\\", "/"))
                if ruta.is_absolute() or ".." in ruta.parts or len(ruta.parts) != 1:
                    errores.append(f"Nombre o ruta insegura dentro del ZIP: {info.filename}")
                    continue
                if ruta.suffix.lower() != ".pdf":
                    errores.append(f"El ZIP contiene un archivo que no es PDF: {ruta.name}")
                    continue
                if info.file_size > MAX_PDF_BYTES:
                    errores.append(f"El documento {ruta.name} excede el límite de 2 MB")
                    continue
                clave = ruta.name.casefold()
                if clave in indice:
                    errores.append(f"Archivo duplicado dentro del ZIP: {ruta.name}")
                    continue
                with archivo_zip.open(info) as stream:
                    if stream.read(5) != b"%PDF-":
                        errores.append(f"El archivo {ruta.name} no contiene un PDF válido")
                        continue
                indice[clave] = info
    except (zipfile.BadZipFile, OSError):
        return {}, ["No fue posible leer el archivo ZIP"]
    return indice, errores


def _tipos_documentales(db: Session) -> dict[str, TipoDocumentoEmpresaModel]:
    tipos = db.query(TipoDocumentoEmpresaModel).filter(TipoDocumentoEmpresaModel.activo.is_(True)).all()
    resultado: dict[str, TipoDocumentoEmpresaModel] = {}
    for tipo in tipos:
        texto = _clave_catalogo(f"{tipo.nombre} {tipo.descripcion or ''}")
        if "fiscal" in texto:
            resultado.setdefault("fiscal", tipo)
        elif "domicilio" in texto:
            resultado.setdefault("domicilio", tipo)
        elif "identificacion" in texto or "representante" in texto:
            resultado.setdefault("identificacion", tipo)
        elif "convenio" in texto or "carta" in texto or "vinculacion" in texto:
            resultado.setdefault("carta", tipo)
        elif "constitut" in texto or "legal" in texto:
            resultado.setdefault("legal", tipo)
    return resultado


def _validar_fila(fila: dict[str, str], numero: int, db: Session) -> dict:
    errores: list[str] = []
    advertencias: list[str] = []
    datos = {columna: _texto(fila.get(columna)) for columna in COLUMNAS}
    if not datos[COLUMNAS[21]]:
        datos[COLUMNAS[21]] = "Solicitante"
    for campo in CAMPOS_OBLIGATORIOS:
        if not datos[campo]:
            errores.append(f"{campo}: campo obligatorio")

    rfc = _normalizar_rfc(datos[COLUMNAS[2]])
    datos[COLUMNAS[2]] = rfc

    for indice, opciones in (
        (17, MODALIDADES), (20, ESTADOS_MEXICO),
        (21, ESTATUS_PERMITIDOS),
    ):
        if not datos[COLUMNAS[indice]]:
            continue
        normal = _catalogo(datos[COLUMNAS[indice]], opciones)
        if normal is None:
            errores.append(f"{COLUMNAS[indice]}: valor fuera del catálogo")
        else:
            datos[COLUMNAS[indice]] = normal

    if datos[COLUMNAS[19]] and not MUNICIPIO_PATTERN.fullmatch(datos[COLUMNAS[19]]):
        errores.append(f"{COLUMNAS[19]}: municipio no válido")

    for indice in (4, 10):
        valor = datos[COLUMNAS[indice]]
        if not valor:
            continue
        telefonos = _normalizar_telefonos(valor)
        if telefonos is None:
            errores.append(f"{COLUMNAS[indice]}: teléfono no válido")
        else:
            datos[COLUMNAS[indice]] = telefonos

    for indice in (5, 11):
        correo = _correo_valido(datos[COLUMNAS[indice]]) if datos[COLUMNAS[indice]] else None
        if datos[COLUMNAS[indice]] and correo is None:
            errores.append(f"{COLUMNAS[indice]}: correo no válido")
        elif correo:
            datos[COLUMNAS[indice]] = correo

    try:
        capacidad_float = float(datos[COLUMNAS[13]])
        if not capacidad_float.is_integer() or capacidad_float <= 0:
            raise ValueError
        datos[COLUMNAS[13]] = str(int(capacidad_float))
    except (TypeError, ValueError):
        if datos[COLUMNAS[13]]:
            errores.append(f"{COLUMNAS[13]}: debe ser un entero mayor que cero")

    if datos[COLUMNAS[7]] and len(datos[COLUMNAS[7]].split()) < 2:
        errores.append(f"{COLUMNAS[7]}: captura nombre y al menos un apellido")
    if datos[COLUMNA_DOCUMENTOS] and not _enlace_documentos_valido(datos[COLUMNA_DOCUMENTOS]):
        errores.append(f"{COLUMNA_DOCUMENTOS.splitlines()[0]}: captura un enlace http:// o https:// válido")
    if datos[COLUMNAS[23]] == "Pendiente":
        advertencias.append("La carta de colaboración quedó pendiente de entrega")
    if datos[COLUMNAS[21]] in {"Pendiente", "Aceptado"}:
        advertencias.append(
            "El estatus capturado se conservará en observaciones; la importación no activa empresas ni crea cuentas de acceso."
        )

    duplicada = db.query(EmpresaModel).filter(func.upper(EmpresaModel.rfc) == rfc).first() if rfc else None
    correo_oficial = datos[COLUMNAS[5]].lower()
    if duplicada is None and correo_oficial:
        duplicada = db.query(EmpresaModel).filter(func.lower(EmpresaModel.correo_contacto) == correo_oficial).first()
    if duplicada is not None:
        errores.append(f"Duplicidad: ya existe {duplicada.nombre_empresa} con ID {duplicada.id_empresa}")

    estado_validacion = "Inválido" if errores else ("Con advertencias" if advertencias else "Válido")
    return {
        "fila": numero, "nombre": datos[COLUMNAS[0]], "rfc": rfc, "tipo_unidad": datos[COLUMNAS[1]],
        "municipio": datos[COLUMNAS[19]], "estado": datos[COLUMNAS[20]], "responsable": datos[COLUMNAS[7]],
        "capacidad": datos[COLUMNAS[13]], "estatus_validacion": estado_validacion, "errores": errores,
        "advertencias": advertencias, "datos": datos, "duplicada": duplicada is not None,
        "id_existente": duplicada.id_empresa if duplicada else None,
    }


def _leer_y_validar(excel: bytes, db: Session) -> tuple[list[dict], list[str]]:
    errores_generales: list[str] = []
    try:
        libro = load_workbook(io.BytesIO(excel), read_only=False, data_only=False, keep_vba=False)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="No fue posible leer el archivo XLSX") from exc
    if HOJA_DATOS not in libro.sheetnames:
        raise HTTPException(status_code=400, detail=f"No existe la hoja '{HOJA_DATOS}'")
    hoja = libro[HOJA_DATOS]

    fila_encabezados = next(
        (
            numero
            for numero in range(1, min(hoja.max_row, 20) + 1)
            if _texto(hoja.cell(numero, 1).value) == COLUMNAS[0]
        ),
        None,
    )
    if fila_encabezados is None:
        raise HTTPException(
            status_code=400,
            detail="No se encontró la fila de encabezados de la plantilla. Descarga una plantilla nueva y no cambies sus títulos.",
        )
    encabezados = [_texto(hoja.cell(fila_encabezados, columna).value) for columna in range(1, len(COLUMNAS) + 1)]
    encabezados_no_vacios = [encabezado for encabezado in encabezados if encabezado]
    if len(encabezados_no_vacios) != len(set(encabezados_no_vacios)):
        errores_generales.append("La hoja contiene encabezados duplicados")
    faltantes = [columna for columna in COLUMNAS if columna not in encabezados]
    if faltantes:
        raise HTTPException(status_code=400, detail="Faltan encabezados: " + ", ".join(faltantes))
    if encabezados != COLUMNAS:
        raise HTTPException(status_code=400, detail="Los encabezados no respetan el orden de la plantilla")

    filas_crudas: list[tuple[int, dict[str, str]]] = []
    filas_vacias: list[int] = []
    encontro_datos = False
    for numero in range(fila_encabezados + 1, hoja.max_row + 1):
        celdas = [hoja.cell(numero, indice + 1) for indice in range(len(COLUMNAS))]
        if any(celda.data_type == "f" for celda in celdas):
            errores_generales.append(f"Fila {numero}: no se permiten fórmulas")
        valores = [_texto(celda.value) for celda in celdas]
        if not any(valores):
            if encontro_datos:
                filas_vacias.append(numero)
            continue
        primer_valor = valores[0].upper()
        if primer_valor.startswith("[EJEMPLO") or primer_valor == "EJEMPLO CORRECTO, NO IMPORTAR":
            continue
        encontro_datos = True
        if filas_vacias:
            errores_generales.append(f"Fila {numero}: existen filas vacías entre registros ({', '.join(map(str, filas_vacias))})")
            filas_vacias.clear()
        filas_crudas.append((numero, dict(zip(COLUMNAS, valores))))
    if len(filas_crudas) > MAX_FILAS:
        raise HTTPException(status_code=400, detail=f"El archivo excede el máximo de {MAX_FILAS} empresas")
    if not filas_crudas:
        raise HTTPException(status_code=400, detail="La plantilla no contiene empresas para validar")

    vistas = [_validar_fila(fila, numero, db) for numero, fila in filas_crudas]
    rfc_filas: dict[str, int] = {}
    correo_filas: dict[str, int] = {}
    archivos_filas: dict[str, int] = {}
    for vista in vistas:
        rfc = vista["rfc"]
        correo = vista["datos"][COLUMNAS[5]].casefold()
        if rfc and rfc in rfc_filas:
            vista["errores"].append(f"RFC duplicado; también aparece en la fila {rfc_filas[rfc]}")
        elif rfc:
            rfc_filas[rfc] = vista["fila"]
        if correo and correo in correo_filas:
            vista["errores"].append(f"Correo institucional duplicado; también aparece en la fila {correo_filas[correo]}")
        elif correo:
            correo_filas[correo] = vista["fila"]
        if vista["errores"]:
            vista["estatus_validacion"] = "Inválido"
    return vistas, errores_generales


def _crear_plantilla(tipos_unidad: list[str]) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = HOJA_DATOS
    catalogos = libro.create_sheet(HOJA_CATALOGOS)
    instrucciones = libro.create_sheet(HOJA_INSTRUCCIONES)

    azul, azul_oscuro, azul_claro = "1565C0", "0D2B5E", "EAF4FF"
    blanco, verde_claro, amarillo, gris = "FFFFFF", "E8F5E9", "FFF7D6", "E7E6E6"
    morado, gris_texto, rojo_claro = "5B3F93", "374151", "FFEBEE"
    borde_fino = Border(
        left=Side(style="thin", color="CBD5E1"), right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"), bottom=Side(style="thin", color="CBD5E1"),
    )
    ultima_columna = get_column_letter(len(COLUMNAS))

    hoja.merge_cells(f"A1:{ultima_columna}1")
    hoja["A1"] = "PLANTILLA DE CARGA MASIVA DE UNIDADES RECEPTORAS"
    hoja["A1"].fill = PatternFill("solid", fgColor=azul)
    hoja["A1"].font = Font(bold=True, color=blanco, size=16)
    hoja["A1"].alignment = Alignment(vertical="center")
    hoja.row_dimensions[1].height = 32

    hoja.merge_cells(f"A3:{ultima_columna}3")
    hoja["A3"] = "INSTRUCCIONES RÁPIDAS"
    hoja["A3"].fill = PatternFill("solid", fgColor=azul_claro)
    hoja["A3"].font = Font(bold=True, color=azul_oscuro, size=12)
    instrucciones_rapidas = [
        (4, "No cambies los nombres ni el orden de las 24 columnas. Captura una empresa por fila.", "Solo el nombre oficial es indispensable; las demás columnas pueden quedar vacías."),
        (5, "Si dejas Estatus vacío se asignará Solicitante. Los datos capturados sí deben tener formato válido.", "En Documentos en formato PDF puedes pegar un enlace http:// o https://; también puede quedar vacío."),
    ]
    for fila, izquierda, derecha in instrucciones_rapidas:
        hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=6)
        hoja.merge_cells(start_row=fila, start_column=7, end_row=fila, end_column=12)
        hoja.cell(fila, 1, izquierda)
        hoja.cell(fila, 7, derecha)
        for celda in (hoja.cell(fila, 1), hoja.cell(fila, 7)):
            celda.fill = PatternFill("solid", fgColor=amarillo)
            celda.font = Font(color=gris_texto)
            celda.alignment = Alignment(wrap_text=True, vertical="top")
        hoja.row_dimensions[fila].height = 36

    hoja.merge_cells(f"A8:{ultima_columna}8")
    hoja["A8"] = "DATOS A CAPTURAR"
    hoja["A8"].fill = PatternFill("solid", fgColor=azul_claro)
    hoja["A8"].font = Font(bold=True, color=azul_oscuro, size=12)

    for indice, encabezado in enumerate(COLUMNAS, 1):
        celda = hoja.cell(FILA_ENCABEZADOS, indice, encabezado)
        es_documento = encabezado == COLUMNA_DOCUMENTOS
        es_obligatorio = encabezado in CAMPOS_OBLIGATORIOS
        color_fondo = morado if es_documento else (azul_oscuro if es_obligatorio else gris)
        color_texto = blanco if es_documento or es_obligatorio else "404040"
        celda.fill = PatternFill("solid", fgColor=color_fondo)
        celda.font = Font(bold=True, color=color_texto, size=9 if es_documento else 10)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = borde_fino
        comentario = "Pega un único enlace http:// o https:// a la carpeta o archivo que reúne los documentos descritos."
        if not es_documento:
            comentario = "Campo obligatorio" if es_obligatorio else "Campo opcional o condicionado"
        celda.comment = Comment(comentario, "Sistema de Prácticas Profesionales")
    hoja.row_dimensions[FILA_ENCABEZADOS].height = 54

    for fila in range(FILA_INICIO_DATOS, FILA_FIN_CAPTURA + 1):
        for columna in range(1, len(COLUMNAS) + 1):
            celda = hoja.cell(fila, columna)
            celda.fill = PatternFill("solid", fgColor=verde_claro)
            celda.border = borde_fino
            celda.alignment = Alignment(wrap_text=True, vertical="top")
        for columna_texto in (3, 5, 11):
            hoja.cell(fila, columna_texto).number_format = "@"
        hoja.cell(fila, 14).number_format = "0"
        hoja.row_dimensions[fila].height = 24

    hoja.merge_cells(f"A{FILA_TITULO_EJEMPLO}:{ultima_columna}{FILA_TITULO_EJEMPLO}")
    hoja.cell(FILA_TITULO_EJEMPLO, 1, "EJEMPLO CORRECTO, NO IMPORTAR")
    hoja.cell(FILA_TITULO_EJEMPLO, 1).fill = PatternFill("solid", fgColor=azul_claro)
    hoja.cell(FILA_TITULO_EJEMPLO, 1).font = Font(bold=True, color=azul_oscuro, size=12)

    ejemplo = [
        "[EJEMPLO - NO IMPORTAR] Empresa Demostración", "Persona jurídica", "XAXX010101000",
        "Avenida Central 100, Colonia Centro, C.P. 29000", "9611234567", "contacto.ejemplo@empresa.mx",
        "Lunes a viernes de 09:00 a 17:00", "María Fernanda López Pérez", "Directora administrativa",
        "Administración", "9617654321", "responsable.ejemplo@empresa.mx", "Sistemas, administración", 5,
        "Estudiantes de informática o administración", "Apoyo en proyectos, documentación y análisis",
        "Lunes a viernes de 09:00 a 14:00", "Presencial",
        "https://drive.google.com/drive/folders/ID_EJEMPLO", "Tuxtla Gutiérrez", "Chiapas", "Solicitante",
        "Fila de ejemplo; el sistema no la procesa", "Sí",
    ]
    for columna, valor in enumerate(ejemplo, 1):
        celda = hoja.cell(FILA_EJEMPLO, columna, valor)
        celda.fill = PatternFill("solid", fgColor=verde_claro)
        celda.border = borde_fino
        celda.font = Font(italic=True, color="666666")
        celda.alignment = Alignment(wrap_text=True, vertical="top")
    hoja.row_dimensions[FILA_EJEMPLO].height = 54

    anchos = [38, 24, 20, 42, 24, 30, 30, 32, 28, 28, 22, 30, 38, 18, 42, 44, 34, 20, 52, 24, 22, 18, 38, 22]
    for indice, ancho in enumerate(anchos, 1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    hoja.freeze_panes = f"A{FILA_INICIO_DATOS}"
    hoja.auto_filter.ref = f"A{FILA_ENCABEZADOS}:{ultima_columna}{FILA_FIN_CAPTURA}"
    hoja.sheet_view.showGridLines = False

    catalogos.append(["Tipo de unidad receptora", "Modalidad", "Estatus", "Carta de colaboración", "Estado"])
    maximo = max(len(tipos_unidad), len(MODALIDADES), len(ESTATUS_PERMITIDOS), len(CARTA_COLABORACION), len(ESTADOS_MEXICO))
    for fila in range(maximo):
        catalogos.append([
            tipos_unidad[fila] if fila < len(tipos_unidad) else "", MODALIDADES[fila] if fila < len(MODALIDADES) else "",
            ESTATUS_PERMITIDOS[fila] if fila < len(ESTATUS_PERMITIDOS) else "",
            CARTA_COLABORACION[fila] if fila < len(CARTA_COLABORACION) else "",
            ESTADOS_MEXICO[fila] if fila < len(ESTADOS_MEXICO) else "",
        ])
    for celda in catalogos[1]:
        celda.fill = PatternFill("solid", fgColor=azul_oscuro)
        celda.font = Font(bold=True, color=blanco)
        celda.border = borde_fino
    for fila in catalogos.iter_rows(min_row=2, max_row=maximo + 1, min_col=1, max_col=5):
        for celda in fila:
            celda.border = borde_fino
    catalogos.freeze_panes = "A2"
    catalogos.auto_filter.ref = "A1:E" + str(maximo + 1)
    catalogos.sheet_view.showGridLines = False
    for columna in "ABCDE":
        catalogos.column_dimensions[columna].width = 32

    validaciones = [
        (f"R{FILA_INICIO_DATOS}:R{FILA_FIN_CAPTURA}", "'" + HOJA_CATALOGOS + "'!$B$2:$B$" + str(len(MODALIDADES) + 1)),
        (f"V{FILA_INICIO_DATOS}:V{FILA_FIN_CAPTURA}", "'" + HOJA_CATALOGOS + "'!$C$2:$C$" + str(len(ESTATUS_PERMITIDOS) + 1)),
        (f"U{FILA_INICIO_DATOS}:U{FILA_FIN_CAPTURA}", "'" + HOJA_CATALOGOS + "'!$E$2:$E$" + str(len(ESTADOS_MEXICO) + 1)),
    ]
    for rango, formula in validaciones:
        validacion = DataValidation(type="list", formula1=formula, allow_blank=False)
        validacion.error = "Selecciona un valor del catálogo."
        validacion.errorTitle = "Valor no permitido"
        validacion.prompt = "Usa uno de los valores de la hoja Catálogos."
        validacion.promptTitle = "Catálogo"
        validacion.showErrorMessage = True
        validacion.showInputMessage = True
        hoja.add_data_validation(validacion)
        validacion.add(rango)

    reglas = [
        "No modificar los encabezados ni eliminar columnas.", "Capturar un registro por fila en el bloque verde.",
        "Solo el nombre oficial es obligatorio; las demás columnas pueden quedar vacías.",
        "La fila de ejemplo no será procesada.", "El RFC, cuando se capture, admite texto libre y se guarda sin espacios ni guiones.",
        "El tipo de unidad receptora, el domicilio y la carta de colaboración admiten texto libre.",
        "Los correos capturados deben tener formato válido.", "Los teléfonos pueden separarse con coma, punto y coma o diagonal.",
        "El número de estudiantes, cuando se capture, debe ser un entero mayor que cero.",
        "En Documentos en formato PDF se puede pegar un enlace http:// o https://; el campo puede quedar vacío.",
        "Si Estatus queda vacío, el sistema asignará Solicitante.",
        "El Excel debe ser .xlsx; no se aceptan macros ni archivos .xlsm.",
        "No repetir RFC ni correos institucionales.", "La importación no crea cuentas ni activa empresas; el estatus capturado se conserva en observaciones para revisión.",
        "Azul: obligatorio. Morado: enlace de documentos. Gris: opcional. Verde: área de captura.",
    ]
    instrucciones.merge_cells("A1:D1")
    instrucciones["A1"] = "INSTRUCCIONES PARA LA CARGA MASIVA DE UNIDADES RECEPTORAS"
    instrucciones["A1"].fill = PatternFill("solid", fgColor=azul)
    instrucciones["A1"].font = Font(bold=True, color=blanco, size=14)
    instrucciones.merge_cells("A3:D3")
    instrucciones["A3"] = "REGLAS DE LLENADO"
    instrucciones["A3"].fill = PatternFill("solid", fgColor=azul_claro)
    instrucciones["A3"].font = Font(bold=True, color=azul_oscuro, size=12)
    for indice, regla in enumerate(reglas, 5):
        instrucciones.cell(indice, 1, f"{indice - 4}.")
        instrucciones.merge_cells(start_row=indice, start_column=2, end_row=indice, end_column=4)
        instrucciones.cell(indice, 2, regla)
        instrucciones.cell(indice, 2).alignment = Alignment(wrap_text=True, vertical="top")
        instrucciones.cell(indice, 2).fill = PatternFill("solid", fgColor=rojo_claro if indice in (5, 6, 13) else blanco)
    instrucciones.column_dimensions["A"].width = 6
    instrucciones.column_dimensions["B"].width = 55
    instrucciones.column_dimensions["C"].width = 30
    instrucciones.column_dimensions["D"].width = 30
    instrucciones.sheet_view.showGridLines = False

    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _guardar_meta(carpeta: Path, meta: dict) -> None:
    (carpeta / "meta.json").write_text(json.dumps(meta, ensure_ascii=False, indent=2), encoding="utf-8")


def _cargar_meta(id_importacion: str, usuario: UsuarioModel) -> tuple[Path, dict]:
    if not re.fullmatch(r"[a-f0-9]{32}", id_importacion):
        raise HTTPException(status_code=404, detail="Importación no encontrada")
    carpeta = IMPORTACIONES_ROOT / id_importacion
    meta_path = carpeta / "meta.json"
    if not meta_path.exists():
        raise HTTPException(status_code=404, detail="Importación no encontrada o expirada")
    try:
        meta = json.loads(meta_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise HTTPException(status_code=404, detail="Importación no disponible") from exc
    if meta.get("id_usuario") != usuario.id_usuario:
        raise HTTPException(status_code=403, detail="No puedes acceder a una importación de otro usuario")
    return carpeta, meta


def _separar_nombre(nombre_completo: str) -> tuple[str, str, str | None]:
    partes = nombre_completo.split()
    if len(partes) == 2:
        return partes[0], partes[1], None
    return " ".join(partes[:-2]), partes[-2], partes[-1]


def _domicilio_y_detalles(datos: dict[str, str]) -> str:
    lineas: list[str] = []
    if datos[COLUMNAS[3]]:
        lineas.append(datos[COLUMNAS[3]])
    detalles = [
        ("Municipio", datos[COLUMNAS[19]]), ("Estado", datos[COLUMNAS[20]]),
        ("Teléfonos institucionales", datos[COLUMNAS[4]]),
        ("Tipo de unidad receptora", datos[COLUMNAS[1]]), ("Horario de atención", datos[COLUMNAS[6]]),
        ("Área del responsable", datos[COLUMNAS[9]]), ("Áreas receptoras", datos[COLUMNAS[12]]),
        ("Capacidad por periodo", datos[COLUMNAS[13]]), ("Perfil académico", datos[COLUMNAS[14]]),
        ("Actividades", datos[COLUMNAS[15]]), ("Horario de prácticas", datos[COLUMNAS[16]]),
        ("Modalidad", datos[COLUMNAS[17]]), ("Carta de colaboración", datos[COLUMNAS[23]]),
    ]
    lineas.extend(f"{etiqueta}: {valor}" for etiqueta, valor in detalles if valor)
    return "\n".join(lineas)


def _observaciones_solicitud(datos: dict[str, str]) -> str:
    lineas = ["Importación masiva."]
    if datos[COLUMNAS[21]]:
        lineas.append(f"Estatus capturado en plantilla: {datos[COLUMNAS[21]]}")
    if datos[COLUMNAS[23]]:
        lineas.append(f"Carta de colaboración capturada: {datos[COLUMNAS[23]]}")
    detalles = [
        ("Tipo", datos[COLUMNAS[1]]), ("Áreas receptoras", datos[COLUMNAS[12]]),
        ("Capacidad por periodo", datos[COLUMNAS[13]]), ("Perfil requerido", datos[COLUMNAS[14]]),
        ("Actividades", datos[COLUMNAS[15]]), ("Horario", datos[COLUMNAS[16]]),
        ("Modalidad", datos[COLUMNAS[17]]), ("Enlace de documentos PDF", datos[COLUMNA_DOCUMENTOS]),
        ("Observaciones", datos[COLUMNAS[22]]),
    ]
    lineas.extend(f"{etiqueta}: {valor}" for etiqueta, valor in detalles if valor)
    return "\n".join(lineas)


def _excel_seguro(valor: object) -> object:
    return "'" + valor if isinstance(valor, str) and valor.startswith(("=", "+", "-", "@")) else valor


def _crear_reporte(resultados: list[dict]) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Resultados"
    encabezados = ["Número de fila original", "Nombre de la empresa", "RFC", "Resultado", "Identificador generado", "Advertencias", "Errores", "Fecha de procesamiento"]
    hoja.append(encabezados)
    for item in resultados:
        hoja.append([
            item["fila"], _excel_seguro(item.get("nombre", "")), _excel_seguro(item.get("rfc", "")),
            item["resultado"], item.get("id_empresa") or item.get("id_existente") or "",
            _excel_seguro("; ".join(item.get("advertencias", []))), _excel_seguro("; ".join(item.get("errores", []))),
            item["fecha_procesamiento"],
        ])
    for celda in hoja[1]:
        celda.fill = PatternFill("solid", fgColor="0D2B5E")
        celda.font = Font(bold=True, color="FFFFFF")
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = "A1:H" + str(max(hoja.max_row, 2))
    hoja.sheet_view.showGridLines = False
    for indice, ancho in enumerate([22, 40, 20, 30, 24, 48, 60, 24], 1):
        hoja.column_dimensions[get_column_letter(indice)].width = ancho
    for fila in hoja.iter_rows(min_row=2):
        for celda in fila:
            celda.alignment = Alignment(vertical="top", wrap_text=True)
    salida = io.BytesIO()
    libro.save(salida)
    return salida.getvalue()


def _guardar_documentos(db: Session, empresa: EmpresaModel, datos: dict[str, str], zip_bytes: bytes | None, tipos: dict[str, TipoDocumentoEmpresaModel]) -> list[Path]:
    if zip_bytes is None:
        return []
    guardados: list[Path] = []
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as archivo_zip:
        infos = {PurePosixPath(item.filename.replace("\\", "/")).name.casefold(): item for item in archivo_zip.infolist() if not item.is_dir()}
        for clave, columna in COLUMNAS_DOCUMENTOS.items():
            nombre = datos[columna]
            tipo = tipos.get(clave)
            if not nombre or tipo is None:
                continue
            info = infos.get(nombre.casefold())
            if info is None:
                raise ValueError(f"No se encontró {nombre} en el ZIP validado")
            contenido = archivo_zip.read(info)
            validar_documento_usuario(contenido, nombre, "application/pdf")
            seguro = normalizar_nombre_archivo(nombre, "documento.pdf")
            carpeta = DOCUMENTOS_ROOT / f"empresa_{empresa.id_empresa}" / f"tipo_{tipo.id_tipo_documento_empresa}"
            carpeta.mkdir(parents=True, exist_ok=True)
            ruta = carpeta / f"{uuid4().hex}_{seguro}"
            ruta.write_bytes(contenido)
            guardados.append(ruta)
            db.add(DocumentoEmpresaModel(
                id_empresa=empresa.id_empresa, id_tipo_documento_empresa=tipo.id_tipo_documento_empresa,
                nombre_archivo=seguro, ruta_archivo=str(ruta), estado_documento="Pendiente",
                observaciones="Cargado mediante importación masiva de unidades receptoras.",
            ))
    return guardados

@router.get("/plantilla")
def descargar_plantilla(db: Session = Depends(obtener_db)):
    tipos_unidad = [
        tipo.nombre
        for tipo in db.query(TipoUnidadReceptoraModel)
        .filter(TipoUnidadReceptoraModel.activo.is_(True))
        .order_by(TipoUnidadReceptoraModel.nombre.asc())
        .all()
    ]
    contenido = _crear_plantilla(tipos_unidad)
    return StreamingResponse(
        io.BytesIO(contenido),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="Plantilla_Carga_Unidades_Receptoras.xlsx"'},
    )


@router.post("/validar")
async def validar_importacion(
    archivo_excel: UploadFile = File(...),
    archivo_zip: UploadFile | None = File(default=None),
    db: Session = Depends(obtener_db),
    usuario_actual: UsuarioModel = Depends(obtener_usuario_actual),
):
    _limpiar_jobs_expirados()
    excel = await archivo_excel.read()
    zip_bytes = await archivo_zip.read() if archivo_zip is not None else None
    _validar_excel_upload(archivo_excel, excel)
    _validar_zip_upload(archivo_zip, zip_bytes)
    vistas, errores_generales = _leer_y_validar(excel, db)
    id_importacion = uuid4().hex
    carpeta = IMPORTACIONES_ROOT / id_importacion
    carpeta.mkdir(parents=True, exist_ok=False)
    (carpeta / "entrada.xlsx").write_bytes(excel)
    if zip_bytes is not None:
        (carpeta / "documentos.zip").write_bytes(zip_bytes)
    meta = {
        "id_importacion": id_importacion, "id_usuario": usuario_actual.id_usuario,
        "fecha_validacion": datetime.now(timezone.utc).isoformat(),
        "archivo_excel": normalizar_nombre_archivo(archivo_excel.filename, "empresas.xlsx"),
        "archivo_zip": normalizar_nombre_archivo(archivo_zip.filename, "documentos.zip") if archivo_zip else None,
        "hash_excel": _sha256(excel), "hash_zip": _sha256(zip_bytes) if zip_bytes is not None else None,
        "errores_generales": errores_generales, "filas": vistas, "procesada": False,
    }
    _guardar_meta(carpeta, meta)
    validas = sum(1 for item in vistas if item["estatus_validacion"] != "Inválido")
    invalidas = len(vistas) - validas
    advertencias = sum(1 for item in vistas if item["estatus_validacion"] == "Con advertencias")
    return {
        "id_importacion": id_importacion,
        "resumen": {"total": len(vistas), "validas": validas, "con_advertencias": advertencias, "invalidas": invalidas},
        "errores_generales": errores_generales,
        "filas": vistas,
        "puede_confirmar": validas > 0 and not errores_generales,
    }


@router.post("/{id_importacion}/confirmar")
def confirmar_importacion(
    id_importacion: str,
    db: Session = Depends(obtener_db),
    usuario_actual: UsuarioModel = Depends(obtener_usuario_actual),
):
    carpeta, meta = _cargar_meta(id_importacion, usuario_actual)
    if meta.get("procesada"):
        return meta["resultado_final"]
    if meta.get("errores_generales"):
        raise HTTPException(status_code=400, detail="La validación contiene errores generales")
    excel_path = carpeta / "entrada.xlsx"
    zip_path = carpeta / "documentos.zip"
    if not excel_path.exists() or _sha256(excel_path.read_bytes()) != meta.get("hash_excel"):
        raise HTTPException(status_code=409, detail="El archivo Excel cambió después de validarse")
    zip_bytes = zip_path.read_bytes() if zip_path.exists() else None
    if zip_bytes is not None and _sha256(zip_bytes) != meta.get("hash_zip"):
        raise HTTPException(status_code=409, detail="El archivo ZIP cambió después de validarse")

    resultados: list[dict] = []
    creadas = omitidas = con_advertencias = con_errores = 0
    for fila in meta["filas"]:
        ahora = datetime.now(timezone.utc).isoformat()
        base = {
            "fila": fila["fila"], "nombre": fila["nombre"], "rfc": fila["rfc"],
            "advertencias": fila.get("advertencias", []), "errores": fila.get("errores", []),
            "fecha_procesamiento": ahora, "id_existente": fila.get("id_existente"),
        }
        if fila["estatus_validacion"] == "Inválido":
            if fila.get("duplicada"):
                base["resultado"] = "Omitida por duplicidad"
                omitidas += 1
            else:
                base["resultado"] = "No importada por errores"
                con_errores += 1
            resultados.append(base)
            continue
        datos = fila["datos"]
        duplicada = None
        if fila["rfc"]:
            duplicada = db.query(EmpresaModel).filter(func.upper(EmpresaModel.rfc) == fila["rfc"]).first()
        if duplicada is None and datos[COLUMNAS[5]]:
            duplicada = db.query(EmpresaModel).filter(func.lower(EmpresaModel.correo_contacto) == datos[COLUMNAS[5]].lower()).first()
        if duplicada is not None:
            base.update({"resultado": "Omitida por duplicidad", "id_existente": duplicada.id_empresa, "errores": [f"La empresa ya existe como {duplicada.nombre_empresa}"]})
            omitidas += 1
            resultados.append(base)
            continue

        rutas_guardadas: list[Path] = []
        try:
            with db.begin_nested():
                tipo_tramite = "Vinculacion" if datos[COLUMNAS[23]] == "No" else "Convenio"
                domicilio = _domicilio_y_detalles(datos) or None
                telefono_principal = datos[COLUMNAS[4]].split("; ", 1)[0] or None
                # La carga masiva registra la solicitud, pero no crea usuario ni valida expediente.
                # Por eso ninguna empresa importada se marca como Activa automáticamente.
                estado_empresa = "Solicitante"
                estado_solicitud = "Recibida"
                id_empresa_reutilizado, registro_identidad = buscar_id_empresa_reutilizable(
                    db,
                    rfc=fila["rfc"] or None,
                    correo=datos[COLUMNAS[5]] or None,
                    nombre=datos[COLUMNAS[0]],
                )
                empresa = EmpresaModel(
                    id_empresa=id_empresa_reutilizado,
                    nombre_empresa=datos[COLUMNAS[0]], rfc=fila["rfc"] or None, giro=datos[COLUMNAS[1]] or None,
                    domicilio=domicilio, telefono=telefono_principal,
                    correo_contacto=datos[COLUMNAS[5]], tipo_tramite=tipo_tramite,
                    estado_empresa=estado_empresa,
                    id_tipo_unidad_receptora=(
                        db.query(TipoUnidadReceptoraModel.id_tipo_unidad_receptora)
                        .filter(
                            TipoUnidadReceptoraModel.nombre == datos[COLUMNAS[1]],
                            TipoUnidadReceptoraModel.activo.is_(True),
                        )
                        .scalar()
                    ),
                )
                db.add(empresa)
                db.flush()
                marcar_id_reutilizado(registro_identidad)
                if datos[COLUMNAS[7]]:
                    nombre, apellido_paterno, apellido_materno = _separar_nombre(datos[COLUMNAS[7]])
                    db.add(ResponsableEmpresaModel(
                        id_empresa=empresa.id_empresa, id_usuario=None, nombre=nombre,
                        apellido_paterno=apellido_paterno, apellido_materno=apellido_materno,
                        cargo=datos[COLUMNAS[8]] or None, telefono=datos[COLUMNAS[10]] or None,
                        correo=datos[COLUMNAS[11]] or None,
                    ))
                db.add(SolicitudEmpresaModel(
                    id_empresa=empresa.id_empresa, tipo_tramite_solicitado=tipo_tramite,
                    estado_solicitud=estado_solicitud, observaciones=_observaciones_solicitud(datos),
                ))
                db.flush()
                id_empresa = empresa.id_empresa
            db.commit()
            base.update({
                "resultado": "Importada con advertencias" if base["advertencias"] else "Importada correctamente",
                "id_empresa": id_empresa,
                "id_reutilizado": id_empresa_reutilizado is not None,
            })
            creadas += 1
            if base["advertencias"]:
                con_advertencias += 1
            registrar_bitacora(
                db, usuario_actual.id_usuario, "Importar empresa desde Excel", "coord_unidades_empresas",
                f"Se importó {fila['nombre']} con RFC {fila['rfc']} desde {meta['archivo_excel']}.", "empresa", id_empresa,
            )
        except Exception as exc:
            db.rollback()
            for ruta in rutas_guardadas:
                ruta.unlink(missing_ok=True)
            base.update({"resultado": "No importada por errores", "errores": [*base["errores"], f"Error al guardar la fila: {str(exc)[:180]}"]})
            con_errores += 1
        resultados.append(base)

    (carpeta / "reporte_resultados.xlsx").write_bytes(_crear_reporte(resultados))
    resumen = {"total": len(resultados), "creadas": creadas, "omitidas": omitidas, "con_advertencias": con_advertencias, "con_errores": con_errores}
    respuesta = {"id_importacion": id_importacion, "resumen": resumen, "resultados": resultados, "reporte_disponible": True}
    meta["procesada"] = True
    meta["fecha_procesamiento"] = datetime.now(timezone.utc).isoformat()
    meta["resultado_final"] = respuesta
    _guardar_meta(carpeta, meta)
    excel_path.unlink(missing_ok=True)
    zip_path.unlink(missing_ok=True)
    registrar_bitacora(
        db, usuario_actual.id_usuario, "Finalizar importación masiva de empresas", "coord_unidades_empresas",
        f"Archivo {meta['archivo_excel']}: {creadas} creadas, {omitidas} omitidas y {con_errores} con errores.",
        "importacion_empresas", None,
    )
    return respuesta


@router.get("/{id_importacion}/reporte")
def descargar_reporte(id_importacion: str, usuario_actual: UsuarioModel = Depends(obtener_usuario_actual)):
    carpeta, meta = _cargar_meta(id_importacion, usuario_actual)
    reporte = carpeta / "reporte_resultados.xlsx"
    if not meta.get("procesada") or not reporte.exists():
        raise HTTPException(status_code=404, detail="El reporte todavía no está disponible")
    return FileResponse(
        reporte,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        filename=f"Reporte_Importacion_Unidades_{id_importacion[:8]}.xlsx",
    )
